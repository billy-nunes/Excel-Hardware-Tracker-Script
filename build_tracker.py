"""
Hardware Request Tracker — Excel workbook generator.

Builds a ready-to-use .xlsx for logging IT hardware requests from ticket
to deployment: dropdowns, auto-calculated cost and age, overdue/late
highlighting, and a summary sheet.

Usage:
    pip install -r requirements.txt
    python build_tracker.py                      # -> Hardware_Request_Tracker.xlsx
    python build_tracker.py -o my_tracker.xlsx   # custom output name
    python build_tracker.py --rows 1000          # pre-format more rows

Note: openpyxl writes formulas but does not calculate them. Excel (or
LibreOffice) calculates everything when the file is opened.
"""

import argparse
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


# ---------------------------------------------------------------------------
# Configuration — edit these to change the tracker's layout and options
# ---------------------------------------------------------------------------

FONT_NAME = "Arial"
HEADER_COLOUR = "1F2937"          # dark grey header background
FORMULA_HEADER_COLOUR = "4B5563"  # lighter grey marks auto-calculated columns
DEFAULT_OVERDUE_DAYS = 14

# (column header, column width)
COLUMNS = [
    ("Ticket Ref", 14),
    ("Date Raised", 13),
    ("Requester", 20),
    ("Store / Dept", 18),
    ("Request Type", 16),
    ("Item Category", 16),
    ("Item Description", 28),
    ("Qty", 6),
    ("Reason / Justification", 30),
    ("Priority", 10),
    ("Approver", 16),
    ("Approval Status", 15),
    ("Approval Date", 13),
    ("Supplier", 16),
    ("PO Number", 13),
    ("Unit Cost (£)", 12),
    ("Total Cost (£)", 13),
    ("Date Ordered", 13),
    ("Expected Delivery", 14),
    ("Date Delivered", 13),
    ("Asset Tag", 13),
    ("Serial Number", 18),
    ("MDM Enrolled", 14),
    ("Status", 17),
    ("Date Closed", 13),
    ("Days Open", 10),
    ("Notes", 34),
]

DATE_COLUMNS = [
    "Date Raised", "Approval Date", "Date Ordered",
    "Expected Delivery", "Date Delivered", "Date Closed",
]
CURRENCY_COLUMNS = ["Unit Cost (£)", "Total Cost (£)"]
WRAP_TEXT_COLUMNS = ["Item Description", "Reason / Justification", "Notes"]
FORMULA_COLUMNS = ["Total Cost (£)", "Days Open"]

# Dropdown options. Each key must match a column header above.
DROPDOWNS = {
    "Request Type": [
        "New starter", "Replacement", "Additional / upgrade", "Repair",
        "Loan", "Leaver return", "Store opening / refit",
    ],
    "Item Category": [
        "Laptop (Windows)", "MacBook", "Desktop", "iPad", "iPhone / mobile",
        "Monitor", "Keyboard / mouse", "Headset", "Docking station", "Printer",
        "Receipt printer", "Barcode scanner", "Cash drawer", "Card terminal",
        "Network kit", "Cables / adapters", "Other",
    ],
    "Priority": ["Critical", "High", "Medium", "Low"],
    "Approval Status": ["Not required", "Pending", "Approved", "Rejected"],
    "MDM Enrolled": ["Meraki", "Jamf", "Intune", "N/A", "Pending"],
    "Status": [
        "New", "Awaiting approval", "Approved", "Ordered", "Delivered",
        "Deployed", "Closed", "On hold", "Rejected / cancelled",
    ],
}

# Background colour for each status value
STATUS_COLOURS = {
    "New": "DBEAFE",
    "Awaiting approval": "FEF3C7",
    "Approved": "E0E7FF",
    "Ordered": "EDE9FE",
    "Delivered": "CCFBF1",
    "Deployed": "DCFCE7",
    "On hold": "F3F4F6",
}

CLOSED_STATUSES = ["Closed", "Rejected / cancelled"]

# One example row so the user can see the expected format
EXAMPLE_ROW = {
    "Ticket Ref": "#SR-12345",
    "Date Raised": date(2026, 10, 14),
    "Requester": "Jane Smith",
    "Store / Dept": "Store 516",
    "Request Type": "Replacement",
    "Item Category": "iPad",
    "Item Description": "iPad 10th gen 64GB Wi-Fi",
    "Qty": 1,
    "Reason / Justification": "Screen cracked, POS till 2 down",
    "Priority": "High",
    "Approver": "Store Ops Manager",
    "Approval Status": "Approved",
    "Approval Date": date(2026, 10, 15),
    "Supplier": "Supplier name",
    "PO Number": "PO-0001",
    "Unit Cost (£)": 349,
    "Date Ordered": date(2026, 10, 15),
    "Expected Delivery": date(2026, 10, 20),
    "Asset Tag": "KG-IPAD-0123",
    "MDM Enrolled": "Meraki",
    "Status": "Ordered",
    "Notes": "Example row — delete or overwrite",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def col(name):
    """Return the Excel column letter for a header name, e.g. 'Status' -> 'X'."""
    headers = [header for header, _ in COLUMNS]
    return get_column_letter(headers.index(name) + 1)


def font(**kwargs):
    return Font(name=FONT_NAME, **kwargs)


def fill(hex_colour):
    return PatternFill("solid", fgColor=hex_colour)


def style_header_row(ws, headers, colour=HEADER_COLOUR):
    for i, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=text)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = fill(colour)
        cell.alignment = Alignment(vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def build_requests_sheet(ws, rows):
    """Main log: headers, formats, formulas, example row, table."""
    last_row = rows + 1  # +1 for the header row
    border = Border(bottom=Side(style="thin", color="D1D5DB"))

    # Headers and widths
    style_header_row(ws, [header for header, _ in COLUMNS])
    for header, width in COLUMNS:
        ws.column_dimensions[col(header)].width = width
    ws.row_dimensions[1].height = 30
    for header in FORMULA_COLUMNS:
        ws[f"{col(header)}1"].fill = fill(FORMULA_HEADER_COLOUR)
        ws[f"{col(header)}1"].comment = Comment(
            "Formula column — auto-calculated, don't overwrite.", "Tracker")

    # Example row (row 2), in grey italics
    for header, value in EXAMPLE_ROW.items():
        ws[f"{col(header)}2"] = value

    # Short aliases keep the formulas readable
    qty, unit = col("Qty"), col("Unit Cost (£)")
    raised, closed = col("Date Raised"), col("Date Closed")

    for r in range(2, last_row + 1):
        # Formulas — blank until the inputs are filled in
        ws[f"{col('Total Cost (£)')}{r}"] = (
            f'=IF(OR({qty}{r}="",{unit}{r}=""),"",{qty}{r}*{unit}{r})')
        ws[f"{col('Days Open')}{r}"] = (
            f'=IF({raised}{r}="","",'
            f'IF({closed}{r}<>"",{closed}{r},TODAY())-{raised}{r})')

        # Base styling for every cell in the row
        for header, _ in COLUMNS:
            cell = ws[f"{col(header)}{r}"]
            cell.font = font(size=10, italic=(r == 2),
                             color="6B7280" if r == 2 else None)
            cell.border = border
            cell.alignment = Alignment(
                vertical="top", wrap_text=header in WRAP_TEXT_COLUMNS)

        # Number formats
        for header in DATE_COLUMNS:
            ws[f"{col(header)}{r}"].number_format = "dd/mm/yyyy"
        for header in CURRENCY_COLUMNS:
            ws[f"{col(header)}{r}"].number_format = '£#,##0.00;(£#,##0.00);"-"'
        ws[f"{col('Days Open')}{r}"].number_format = "0"

    # Freeze header row + first two columns, and make it an Excel table
    ws.freeze_panes = "C2"
    table = Table(displayName="HardwareRequests",
                  ref=f"A1:{col(COLUMNS[-1][0])}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1",
                                          showRowStripes=True)
    ws.add_table(table)


def build_lists_sheet(ws):
    """Dropdown source lists + the overdue threshold setting."""
    style_header_row(ws, list(DROPDOWNS.keys()))
    for i, values in enumerate(DROPDOWNS.values(), start=1):
        ws.column_dimensions[get_column_letter(i)].width = 24
        for j, value in enumerate(values, start=2):
            ws.cell(row=j, column=i, value=value).font = font()

    # Settings block in H/I
    ws["H1"] = "Settings"
    ws["H1"].font = font(bold=True)
    ws["H2"] = "Overdue after (days)"
    ws["H2"].font = font()
    ws["I2"] = DEFAULT_OVERDUE_DAYS
    ws["I2"].font = font(color="0000FF")  # blue = editable input
    ws["I2"].fill = fill("FFFF00")
    ws["I2"].comment = Comment(
        "Open requests older than this are highlighted red. "
        "Change to suit your SLA.", "Tracker")
    ws["H3"] = ("Lists feed the dropdowns on Requests — edit items in place "
                "(extend the range in Data > Validation if you add rows).")
    ws["H3"].font = font(italic=True, size=9, color="6B7280")
    ws.column_dimensions["H"].width = 22


def add_validation(requests_ws, rows):
    """Dropdowns on the Requests sheet, plus date-only checks."""
    last_row = rows + 1

    for i, (header, values) in enumerate(DROPDOWNS.items(), start=1):
        list_col = get_column_letter(i)
        source = f"Lists!${list_col}$2:${list_col}${len(values) + 1}"
        dv = DataValidation(type="list", formula1=source, allow_blank=True,
                            showErrorMessage=False)  # allow free text too
        requests_ws.add_data_validation(dv)
        dv.add(f"{col(header)}2:{col(header)}{last_row}")

    date_check = DataValidation(
        type="date", operator="greaterThan", formula1="DATE(2020,1,1)",
        allow_blank=True, showErrorMessage=True,
        error="Enter a date (dd/mm/yyyy)")
    requests_ws.add_data_validation(date_check)
    for header in DATE_COLUMNS:
        date_check.add(f"{col(header)}2:{col(header)}{last_row}")


def add_conditional_formatting(ws, rows):
    """Colour rules: closed rows grey, overdue red, late delivery orange."""
    last_row = rows + 1
    status, days = col("Status"), col("Days Open")
    expected, delivered = col("Expected Delivery"), col("Date Delivered")

    # Formulas are written for row 2; Excel adjusts them for each row.
    is_closed = "OR(" + ",".join(f'${status}2="{s}"' for s in CLOSED_STATUSES) + ")"

    whole_table = f"A2:{col(COLUMNS[-1][0])}{last_row}"
    ws.conditional_formatting.add(whole_table, FormulaRule(
        formula=[f'AND(${status}2<>"",{is_closed})'],
        font=Font(color="9CA3AF")))

    ws.conditional_formatting.add(f"{days}2:{days}{last_row}", FormulaRule(
        formula=[f'AND(${days}2<>"",NOT({is_closed}),${days}2>Lists!$I$2)'],
        fill=fill("FECACA"), font=Font(bold=True, color="991B1B")))

    ws.conditional_formatting.add(f"{expected}2:{expected}{last_row}", FormulaRule(
        formula=[f'AND(${expected}2<>"",${delivered}2="",'
                 f'${expected}2<TODAY(),NOT({is_closed}))'],
        fill=fill("FED7AA")))

    for value, colour in STATUS_COLOURS.items():
        ws.conditional_formatting.add(f"{status}2:{status}{last_row}", FormulaRule(
            formula=[f'${status}2="{value}"'], fill=fill(colour)))


def build_summary_sheet(ws, rows):
    """Headline counts and breakdowns, all driven by formulas."""
    last_row = rows + 1

    def rng(header):
        c = col(header)
        return f"Requests!${c}$2:${c}${last_row}"

    status_rng, days_rng = rng("Status"), rng("Days Open")
    not_closed = "*".join(f'({status_rng}<>"{s}")' for s in CLOSED_STATUSES)
    closed_count = "-".join(f'COUNTIF({status_rng},"{s}")' for s in CLOSED_STATUSES)

    ws["A1"] = "Hardware Requests — Summary"
    ws["A1"].font = font(bold=True, size=14)

    metrics = [
        ("Total requests",
         f"=COUNTA({rng('Date Raised')})", "0"),
        ("Open (not closed/rejected)",
         f"=COUNTA({rng('Date Raised')})-{closed_count}", "0"),
        ("Overdue (open > threshold)",
         f'=SUMPRODUCT(({days_rng}<>"")*({days_rng}>Lists!$I$2)*{not_closed})', "0"),
        ("Awaiting approval",
         f'=COUNTIF({status_rng},"Awaiting approval")', "0"),
        ("Late deliveries",
         f'=SUMPRODUCT(({rng("Expected Delivery")}<>"")'
         f'*({rng("Date Delivered")}="")'
         f'*({rng("Expected Delivery")}<TODAY())*{not_closed})', "0"),
        ("Avg days open (open requests)",
         f'=IFERROR(SUMPRODUCT(({days_rng}<>"")*{not_closed},{days_rng})/B4,0)', "0.0"),
        ("Total spend (£, excl. rejected)",
         f'=SUMIFS({rng("Total Cost (£)")},{status_rng},"<>Rejected / cancelled")',
         "£#,##0.00"),
    ]
    for i, (label, formula, number_format) in enumerate(metrics, start=3):
        ws.cell(row=i, column=1, value=label).font = font()
        cell = ws.cell(row=i, column=2, value=formula)
        cell.font = font(bold=True)
        cell.number_format = number_format

    # Breakdown tables
    def breakdown(title, values, source, label_col, count_col):
        ws[f"{label_col}11"] = title
        ws[f"{label_col}11"].font = font(bold=True)
        for i, value in enumerate(values, start=12):
            ws[f"{label_col}{i}"] = value
            ws[f"{label_col}{i}"].font = font()
            ws[f"{count_col}{i}"] = f"=COUNTIF({source},{label_col}{i})"
            ws[f"{count_col}{i}"].font = font()

    breakdown("By status", DROPDOWNS["Status"], status_rng, "A", "B")
    breakdown("By category", DROPDOWNS["Item Category"],
              rng("Item Category"), "D", "E")

    for letter, width in {"A": 32, "B": 14, "D": 22, "E": 10}.items():
        ws.column_dimensions[letter].width = width
    ws["A30"] = ("Note: the example row on Requests (row 2) is counted "
                 "until you overwrite or delete it.")
    ws["A30"].font = font(italic=True, size=9, color="6B7280")


def build_guide_sheet(ws):
    lines = [
        "How to use",
        "• Add one row per request on Requests. Ticket ref + Date Raised are the minimum.",
        "• Dropdowns: Request Type, Item Category, Priority, Approval Status, "
        "MDM Enrolled, Status (edit options on Lists).",
        "• Grey-headed columns are formulas — don't type in them: Total Cost = Qty × "
        "Unit Cost; Days Open = today (or Date Closed) − Date Raised.",
        "• Days Open turns red when an open request passes the threshold on "
        f"Lists!I2 (default {DEFAULT_OVERDUE_DAYS} days).",
        "• Expected Delivery turns orange if the date has passed with nothing "
        "in Date Delivered.",
        "• Closed and Rejected rows grey out. Set Date Closed when you close so "
        "Days Open stops counting.",
        "• Record Asset Tag + Serial at deployment so the request ties to the "
        "asset / MDM record.",
        "• Row 2 is an example — overwrite or delete it.",
        "• Use the header filters to slice by store, status or requester.",
    ]
    for i, text in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.font = font(bold=(i == 1), size=14 if i == 1 else 10)
    ws.column_dimensions["A"].width = 110


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def build_workbook(rows):
    wb = Workbook()
    requests_ws = wb.active
    requests_ws.title = "Requests"
    summary_ws = wb.create_sheet("Summary")
    lists_ws = wb.create_sheet("Lists")
    guide_ws = wb.create_sheet("How to use")

    build_requests_sheet(requests_ws, rows)
    build_lists_sheet(lists_ws)
    add_validation(requests_ws, rows)
    add_conditional_formatting(requests_ws, rows)
    build_summary_sheet(summary_ws, rows)
    build_guide_sheet(guide_ws)

    wb.active = 0
    return wb


def main():
    parser = argparse.ArgumentParser(description="Generate the hardware request tracker.")
    parser.add_argument("-o", "--output", default="Hardware_Request_Tracker.xlsx",
                        help="output file name (default: %(default)s)")
    parser.add_argument("--rows", type=int, default=500,
                        help="number of pre-formatted data rows (default: %(default)s)")
    args = parser.parse_args()

    build_workbook(args.rows).save(args.output)
    print(f"Saved {args.output} ({args.rows} rows)")


if __name__ == "__main__":
    main()
